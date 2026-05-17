import streamlit as st
import pandas as pd
import openpyxl
from openpyxl import load_workbook
import requests
import json
import io
import copy

# ── Page config ──────────────────────────────────────────────────────────────
st.set_page_config(
    page_title="OMAC AI Excel Data Entry Tool",
    page_icon="📊",
    layout="wide",
    initial_sidebar_state="collapsed",
)

# ── Constants ─────────────────────────────────────────────────────────────────
GROQ_API_KEY = "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"
GROQ_URL = "https://api.groq.com/openai/v1/chat/completions"
GROQ_MODEL = "llama3-70b-8192"
BATCH_SIZE = 5

# ── Custom CSS ────────────────────────────────────────────────────────────────
st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700&display=swap');

html, body, [class*="css"] { font-family: 'Inter', sans-serif; }

/* Header */
.omac-header {
    background: linear-gradient(135deg, #1e3a8a 0%, #3b5bdb 100%);
    border-radius: 14px;
    padding: 22px 28px;
    margin-bottom: 24px;
    display: flex;
    align-items: center;
    justify-content: space-between;
}
.omac-title { color: #fff; font-size: 22px; font-weight: 700; margin: 0; }
.omac-sub   { color: #bfdbfe; font-size: 13px; margin: 2px 0 0; }

/* Step badge */
.step-badge {
    display: inline-block;
    background: #3b5bdb;
    color: #fff;
    border-radius: 50%;
    width: 26px; height: 26px;
    line-height: 26px;
    text-align: center;
    font-weight: 700;
    font-size: 13px;
    margin-right: 8px;
}
.step-title { font-weight: 600; font-size: 15px; }

/* Cards */
.info-card {
    background: #eff6ff;
    border: 1px solid #bfdbfe;
    border-radius: 10px;
    padding: 12px 16px;
    font-size: 13px;
    color: #1e40af;
    margin-bottom: 8px;
}
.success-card {
    background: #f0fdf4;
    border: 1px solid #86efac;
    border-radius: 10px;
    padding: 14px 18px;
    color: #166534;
    font-size: 14px;
}
.warn-card {
    background: #fffbeb;
    border: 1px solid #fcd34d;
    border-radius: 10px;
    padding: 12px 16px;
    font-size: 13px;
    color: #92400e;
}

/* Footer */
.omac-footer {
    text-align: center;
    color: #94a3b8;
    font-size: 12px;
    margin-top: 40px;
    padding-top: 16px;
    border-top: 1px solid #e2e8f0;
}
.omac-footer strong { color: #3b5bdb; }

/* Log box */
.log-box {
    background: #0f172a;
    border-radius: 10px;
    padding: 14px 16px;
    font-family: monospace;
    font-size: 12px;
    max-height: 260px;
    overflow-y: auto;
}
.log-success { color: #4ade80; }
.log-error   { color: #f87171; }
.log-info    { color: #94a3b8; }
</style>
""", unsafe_allow_html=True)

# ── Header ────────────────────────────────────────────────────────────────────
st.markdown("""
<div class="omac-header">
  <div>
    <div class="omac-title">📊 OMAC AI Excel Data Entry Tool</div>
    <div class="omac-sub">AI-powered smart fill for your Excel workbooks</div>
  </div>
  <div style="color:#bfdbfe; font-size:13px; text-align:right;">
    Powered by Groq LLaMA 3<br>
    <span style="color:#93c5fd;">llama3-70b-8192</span>
  </div>
</div>
""", unsafe_allow_html=True)

# ── Groq helper ───────────────────────────────────────────────────────────────
def call_groq(messages):
    headers = {
        "Content-Type": "application/json",
        "Authorization": f"Bearer {GROQ_API_KEY}",
    }
    payload = {
        "model": GROQ_MODEL,
        "messages": messages,
        "max_tokens": 4096,
        "temperature": 0.2,
    }
    resp = requests.post(GROQ_URL, headers=headers, json=payload, timeout=60)
    resp.raise_for_status()
    return resp.json()["choices"][0]["message"]["content"]


def parse_json_safe(text):
    try:
        clean = text.replace("```json", "").replace("```", "").strip()
        return json.loads(clean)
    except Exception:
        return None


# ── Session state defaults ────────────────────────────────────────────────────
for k, v in {
    "wb_bytes": None, "file_name": "", "sheet_name": "",
    "df": None, "result_df": None, "log": [],
}.items():
    if k not in st.session_state:
        st.session_state[k] = v

# ═════════════════════════════════════════════════════════════════════════════
# STEP 1 – Upload
# ═════════════════════════════════════════════════════════════════════════════
st.markdown('<span class="step-badge">1</span><span class="step-title">Upload Workbook</span>', unsafe_allow_html=True)

uploaded = st.file_uploader(
    "Drop your Excel file here",
    type=["xlsx", "xls", "csv"],
    label_visibility="collapsed",
)

if uploaded:
    st.session_state.wb_bytes = uploaded.read()
    st.session_state.file_name = uploaded.name
    st.session_state.result_df = None
    st.session_state.log = []
    st.markdown(f'<div class="info-card">📂 <b>{uploaded.name}</b> loaded successfully</div>', unsafe_allow_html=True)

if not st.session_state.wb_bytes:
    st.markdown('<div class="warn-card">⬆ Please upload an Excel workbook to begin.</div>', unsafe_allow_html=True)
    st.markdown('<div class="omac-footer"><strong>OMAC AI Excel Data Entry Tool</strong> · Developed by <strong>S M Baqir</strong></div>', unsafe_allow_html=True)
    st.stop()

# ═════════════════════════════════════════════════════════════════════════════
# STEP 2 – Sheet selection
# ═════════════════════════════════════════════════════════════════════════════
st.markdown("---")
st.markdown('<span class="step-badge">2</span><span class="step-title">Select Sheet</span>', unsafe_allow_html=True)

fname = st.session_state.file_name.lower()
wb_bytes = st.session_state.wb_bytes

if fname.endswith(".csv"):
    sheet_names = ["Sheet1"]
    sheet_name = "Sheet1"
else:
    wb_peek = load_workbook(io.BytesIO(wb_bytes), read_only=True, data_only=True)
    sheet_names = wb_peek.sheetnames
    wb_peek.close()
    sheet_name = st.selectbox("Choose a sheet", sheet_names)

st.session_state.sheet_name = sheet_name

# Load into DataFrame
if fname.endswith(".csv"):
    df = pd.read_csv(io.BytesIO(wb_bytes), dtype=str).fillna("")
else:
    df = pd.read_excel(io.BytesIO(wb_bytes), sheet_name=sheet_name, dtype=str).fillna("")

st.session_state.df = df
st.caption(f"📋  {len(df)} rows · {len(df.columns)} columns")

# ═════════════════════════════════════════════════════════════════════════════
# STEP 3 – Column & Instruction
# ═════════════════════════════════════════════════════════════════════════════
st.markdown("---")
st.markdown('<span class="step-badge">3</span><span class="step-title">Configure AI Fill</span>', unsafe_allow_html=True)

col1, col2 = st.columns([1, 1])

with col1:
    target_col = st.selectbox("Column to fill (blank cells only)", df.columns.tolist())

with col2:
    empty_count = (df[target_col] == "").sum()
    filled_count = len(df) - empty_count
    st.metric("Empty cells to fill", empty_count)
    st.caption(f"{filled_count} cells already have data — they will NOT be changed")

instruction = st.text_area(
    "Instructions for AI",
    height=130,
    placeholder=(
        "Examples:\n"
        "• Fill based on the complaint description — classify as: Service Delay / Quality of Care / Staff Attitude\n"
        "• If 'Departmental Response Received' is filled, write 'Received', else 'Pending'\n"
        "• Derive from 'Tasks/Problems' column — generate a short action plan\n"
        "• Use correlation between Department and Affinity columns to suggest category"
    ),
)

# Reference columns hint
other_cols = [c for c in df.columns if c != target_col]
with st.expander("📎 Available reference columns for AI context"):
    st.write(", ".join(other_cols))

# ═════════════════════════════════════════════════════════════════════════════
# STEP 4 – Run
# ═════════════════════════════════════════════════════════════════════════════
st.markdown("---")
st.markdown('<span class="step-badge">4</span><span class="step-title">Run AI Fill</span>', unsafe_allow_html=True)

run_col, dl_col = st.columns([2, 1])

with run_col:
    run_btn = st.button(
        "▶  Run AI Fill",
        type="primary",
        disabled=(not instruction or empty_count == 0),
        use_container_width=True,
    )

if run_btn:
    st.session_state.log = []

    # ── Reset DataFrame index so it is always 0-based ──────────────────────
    result_df = df.copy().reset_index(drop=True)
    work_df   = df.copy().reset_index(drop=True)

    # ── Detect ALL blank cells: empty string, NaN, None, whitespace-only ───
    def is_blank(v):
        if v is None:
            return True
        return str(v).strip() == ""

    empty_rows = [i for i in range(len(work_df)) if is_blank(work_df.at[i, target_col])]

    progress_bar   = st.progress(0, text="Starting…")
    log_placeholder = st.empty()
    logs = []

    def render_log():
        html = '<div class="log-box">'
        for l in logs[-40:]:
            cls  = {"success": "log-success", "error": "log-error"}.get(l["t"], "log-info")
            icon = {"success": "✓", "error": "✗"}.get(l["t"], "›")
            html += f'<div class="{cls}">{icon} {l["m"]}</div>'
        html += "</div>"
        log_placeholder.markdown(html, unsafe_allow_html=True)

    logs.append({"m": f"Found {len(empty_rows)} empty cells in '{target_col}'", "t": "info"})
    render_log()

    total = len(empty_rows)
    done  = 0

    for b in range(0, total, BATCH_SIZE):
        chunk_idx = empty_rows[b : b + BATCH_SIZE]   # list of 0-based df row indices
        rows_payload = []
        for ri in chunk_idx:
            row_dict = {col: str(work_df.at[ri, col]) for col in work_df.columns}
            # Use a stable key = the 0-based index so we can map back exactly
            rows_payload.append({"_rowIdx": ri, "rowNumber": ri + 2, "data": row_dict})

        system_prompt = (
            f'You are an expert data analyst filling missing Excel values.\n'
            f'Target column: "{target_col}"\n'
            f'User instruction: {instruction}\n'
            f'All column headers: {", ".join(work_df.columns.tolist())}\n'
            f'Rules:\n'
            f'- Return ONLY a valid JSON array, no explanation, no markdown fences\n'
            f'- Fill a value for EVERY row in the input — do not skip any\n'
            f'- Match the format/style of existing values in the column if any exist\n'
            f'- If no prior values exist, follow the instruction strictly\n'
            f'- Keep the "_rowIdx" field exactly as given so values map to the right row\n'
            f'- Format: [{{"_rowIdx": N, "value": "..."}}]'
        )
        user_msg = (
            f'Fill "{target_col}" for ALL {len(chunk_idx)} rows below.\n'
            f'Return exactly {len(chunk_idx)} items in the JSON array.\n'
            + json.dumps(rows_payload, ensure_ascii=False, indent=2)
        )

        try:
            logs.append({"m": f"Batch {b//BATCH_SIZE+1}: rows {chunk_idx[0]+2}–{chunk_idx[-1]+2}  ({len(chunk_idx)} cells)…", "t": "info"})
            render_log()

            response = call_groq([
                {"role": "system", "content": system_prompt},
                {"role": "user",   "content": user_msg},
            ])
            parsed = parse_json_safe(response)

            if parsed and isinstance(parsed, list):
                filled_indices = set()
                for item in parsed:
                    # Accept either _rowIdx (preferred) or rowNumber fallback
                    if "_rowIdx" in item:
                        ri = int(item["_rowIdx"])
                    else:
                        ri = int(item.get("rowNumber", -1)) - 2
                    val = str(item.get("value", "")).strip()
                    if 0 <= ri < len(result_df) and ri in chunk_idx:
                        result_df.at[ri, target_col] = val
                        filled_indices.add(ri)
                        logs.append({"m": f"Row {ri+2}: {val[:90]}", "t": "success"})

                # If AI skipped some rows in this batch, retry row-by-row
                skipped = [ri for ri in chunk_idx if ri not in filled_indices]
                for ri in skipped:
                    logs.append({"m": f"Retrying row {ri+2} individually…", "t": "info"})
                    render_log()
                    row_dict = {col: str(work_df.at[ri, col]) for col in work_df.columns}
                    retry_msg = (
                        f'Fill the single missing value for column "{target_col}".\n'
                        f'Instruction: {instruction}\n'
                        f'Row data: {json.dumps(row_dict, ensure_ascii=False)}\n'
                        f'Return ONLY: [{{"_rowIdx": {ri}, "value": "..."}}]'
                    )
                    try:
                        r2 = call_groq([
                            {"role": "system", "content": system_prompt},
                            {"role": "user",   "content": retry_msg},
                        ])
                        p2 = parse_json_safe(r2)
                        if p2 and isinstance(p2, list) and p2:
                            val = str(p2[0].get("value", "")).strip()
                            result_df.at[ri, target_col] = val
                            logs.append({"m": f"Row {ri+2} (retry): {val[:90]}", "t": "success"})
                        else:
                            logs.append({"m": f"Row {ri+2}: could not fill after retry", "t": "error"})
                    except Exception as e2:
                        logs.append({"m": f"Row {ri+2} retry error: {str(e2)}", "t": "error"})

            else:
                # Entire batch parse failed — retry each row individually
                logs.append({"m": f"Batch {b//BATCH_SIZE+1} parse failed — retrying row by row…", "t": "error"})
                for ri in chunk_idx:
                    row_dict = {col: str(work_df.at[ri, col]) for col in work_df.columns}
                    retry_msg = (
                        f'Fill the single missing value for column "{target_col}".\n'
                        f'Instruction: {instruction}\n'
                        f'Row data: {json.dumps(row_dict, ensure_ascii=False)}\n'
                        f'Return ONLY: [{{"_rowIdx": {ri}, "value": "..."}}]'
                    )
                    try:
                        r2 = call_groq([
                            {"role": "system", "content": system_prompt},
                            {"role": "user",   "content": retry_msg},
                        ])
                        p2 = parse_json_safe(r2)
                        if p2 and isinstance(p2, list) and p2:
                            val = str(p2[0].get("value", "")).strip()
                            result_df.at[ri, target_col] = val
                            logs.append({"m": f"Row {ri+2}: {val[:90]}", "t": "success"})
                        else:
                            logs.append({"m": f"Row {ri+2}: skipped (no parseable response)", "t": "error"})
                    except Exception as e2:
                        logs.append({"m": f"Row {ri+2} error: {str(e2)}", "t": "error"})

        except Exception as e:
            logs.append({"m": f"Batch error: {str(e)}", "t": "error"})

        done += len(chunk_idx)
        pct = done / total
        progress_bar.progress(pct, text=f"{done}/{total} cells processed ({int(pct*100)}%)")
        render_log()

    progress_bar.progress(1.0, text="✓ Complete!")
    logs.append({"m": "All done!", "t": "success"})
    render_log()
    st.session_state.result_df = result_df
    st.session_state.log = logs

# ── Download ──────────────────────────────────────────────────────────────────
if st.session_state.result_df is not None:
    result_df = st.session_state.result_df

    st.markdown('<div class="success-card">✓ AI fill complete! Download your workbook below — only blank cells were changed.</div>', unsafe_allow_html=True)
    st.markdown("")

    # Build output preserving original workbook structure
    out = io.BytesIO()

    if fname.endswith(".csv"):
        result_df.to_csv(out, index=False)
        out.seek(0)
        dl_name = st.session_state.file_name.replace(".csv", "_AI_filled.csv")
        mime = "text/csv"
    else:
        # Load original workbook to preserve styles/merges/other sheets
        orig_wb = load_workbook(io.BytesIO(st.session_state.wb_bytes))
        ws = orig_wb[st.session_state.sheet_name]

        # Find header row (row 1) and build col index map
        headers_in_ws = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        col_map = {str(h): i + 1 for i, h in enumerate(headers_in_ws) if h is not None}

        # Only write to cells that were originally empty
        orig_df = st.session_state.df.reset_index(drop=True)
        for ri, row_series in result_df.iterrows():
            orig_val = orig_df.at[ri, target_col]
            new_val  = row_series[target_col]
            if (orig_val is None or str(orig_val).strip() == "") and str(new_val).strip() != "":
                excel_row = ri + 2  # +1 for header, +1 for 1-based
                excel_col = col_map.get(target_col)
                if excel_col:
                    ws.cell(row=excel_row, column=excel_col).value = new_val

        orig_wb.save(out)
        out.seek(0)
        base = st.session_state.file_name.rsplit(".", 1)[0]
        dl_name = f"{base}_AI_filled.xlsx"
        mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    st.download_button(
        label="⬇  Download Filled Workbook",
        data=out,
        file_name=dl_name,
        mime=mime,
        type="primary",
        use_container_width=True,
    )

    # Preview
    st.markdown("---")
    st.markdown("**Preview — filled data (first 10 rows)**")
    preview_cols = [target_col] + [c for c in result_df.columns if c != target_col][:5]
    st.dataframe(
        result_df[preview_cols].head(10),
        use_container_width=True,
        height=320,
    )

# ── Footer ────────────────────────────────────────────────────────────────────
st.markdown(
    '<div class="omac-footer"><strong>OMAC AI Excel Data Entry Tool</strong> &nbsp;·&nbsp; Developed by <strong>S M Baqir</strong></div>',
    unsafe_allow_html=True,
)
